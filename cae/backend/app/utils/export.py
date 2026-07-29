"""Excel and PDF export utilities."""
from __future__ import annotations

import io
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from sqlalchemy.orm import Session


def _flatten_evidence(data) -> str:
    if not data:
        return ""
    if isinstance(data, dict):
        return " | ".join(f"{k}: {v}" for k, v in list(data.items())[:8])
    if isinstance(data, list):
        return "; ".join(str(item)[:60] for item in data[:5])
    return str(data)[:200]


def generate_working_paper_excel(db: "Session", company_id: str, period: str | None = None) -> io.BytesIO:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from app.models.audit import AuditException
    from app.models.tenant import Company
    import uuid

    wb = openpyxl.Workbook()

    company = db.query(Company).filter(Company.id == uuid.UUID(company_id)).first()
    company_name = company.name if company else "Company"
    company_gstin = getattr(company, 'gstin', '') or ''

    # ── Cover sheet ──────────────────────────────────────────────────────────
    ws_cover = wb.active
    ws_cover.title = "Cover"
    ws_cover.column_dimensions["A"].width = 30
    ws_cover.column_dimensions["B"].width = 40

    def cover_row(ws, row, label, value, bold_val=False):
        ws.cell(row=row, column=1, value=label).font = Font(bold=True, color="888888", size=10)
        c = ws.cell(row=row, column=2, value=value)
        if bold_val:
            c.font = Font(bold=True, size=11)

    ws_cover["A1"] = "INTERNAL AUDIT WORKING PAPER"
    ws_cover["A1"].font = Font(bold=True, size=18, color="C8A951")
    ws_cover.merge_cells("A1:B1")
    ws_cover.row_dimensions[1].height = 32

    cover_row(ws_cover, 3, "Company", company_name, bold_val=True)
    cover_row(ws_cover, 4, "GSTIN", company_gstin)
    cover_row(ws_cover, 5, "Period", period or "All Periods")
    cover_row(ws_cover, 6, "Generated", __import__('datetime').datetime.now().strftime('%d %b %Y %H:%M'))

    q_all = db.query(AuditException).filter(AuditException.company_id == uuid.UUID(company_id))
    total = q_all.count()
    cover_row(ws_cover, 8, "Total Exceptions", total, bold_val=True)

    # ── Exceptions sheet ─────────────────────────────────────────────────────
    ws = wb.create_sheet("Exceptions")
    headers = [
        "Exc #", "Category", "Title", "Severity", "Risk Impact",
        "Financial Impact (₹)", "Status", "Detected On",
        "Description", "Evidence Summary", "AI Analysis",
    ]
    header_fill = PatternFill(start_color="0D1117", end_color="0D1117", fill_type="solid")
    header_font = Font(bold=True, color="C8A951", size=10)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    severity_colors = {"critical": "FF0000", "high": "FF6B00", "medium": "FFC107", "low": "28A745"}

    exceptions = q_all.order_by(AuditException.detected_on.desc()).all()
    total_impact = 0.0

    for i, exc in enumerate(exceptions, 2):
        impact = float(exc.financial_impact or 0)
        total_impact += impact

        ws.cell(row=i, column=1, value=i - 1)
        ws.cell(row=i, column=2, value=exc.category or "")
        ws.cell(row=i, column=3, value=exc.title or "")

        sev_cell = ws.cell(row=i, column=4, value=(exc.severity or "").upper())
        color = severity_colors.get((exc.severity or "").lower(), "CCCCCC")
        sev_cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        sev_cell.font = Font(bold=True, color="FFFFFF", size=9)
        sev_cell.alignment = Alignment(horizontal="center")

        ws.cell(row=i, column=5, value=exc.risk_impact or "")
        ws.cell(row=i, column=6, value=impact)
        ws.cell(row=i, column=7, value=exc.status or "")
        ws.cell(row=i, column=8, value=str(exc.detected_on) if exc.detected_on else "")

        desc_cell = ws.cell(row=i, column=9, value=exc.description or "")
        desc_cell.alignment = Alignment(wrap_text=True, vertical="top")

        ev_cell = ws.cell(row=i, column=10, value=_flatten_evidence(exc.supporting_data))
        ev_cell.alignment = Alignment(wrap_text=True, vertical="top")

        ai_cell = ws.cell(row=i, column=11, value=exc.ai_analysis or "")
        ai_cell.alignment = Alignment(wrap_text=True, vertical="top")

        ws.row_dimensions[i].height = 60

    # Totals row
    total_row = len(exceptions) + 2
    ws.cell(row=total_row, column=5, value="TOTAL FINANCIAL IMPACT").font = Font(bold=True)
    total_cell = ws.cell(row=total_row, column=6, value=total_impact)
    total_cell.font = Font(bold=True, color="C8A951")

    # Column widths
    col_widths = [6, 18, 45, 11, 20, 22, 16, 14, 50, 45, 50]
    for col, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generate_exception_summary_excel(db: "Session", company_id: str) -> io.BytesIO:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from app.models.audit import AuditException
    from app.models.tenant import Company
    from sqlalchemy import func
    import uuid

    wb = openpyxl.Workbook()
    company = db.query(Company).filter(Company.id == uuid.UUID(company_id)).first()
    company_name = company.name if company else "Company"
    cid = uuid.UUID(company_id)

    header_fill = PatternFill(start_color="0D1117", end_color="0D1117", fill_type="solid")
    header_font = Font(bold=True, color="C8A951", size=10)
    severity_colors = {"critical": "FF0000", "high": "FF6B00", "medium": "FFC107", "low": "28A745"}

    # ── Sheet 1: Overview ─────────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Overview"
    ws1["A1"] = f"Exception Summary — {company_name}"
    ws1["A1"].font = Font(bold=True, size=14, color="C8A951")
    ws1.merge_cells("A1:D1")
    ws1.row_dimensions[1].height = 24

    ws1.cell(row=3, column=1, value="BY SEVERITY").font = Font(bold=True, color="888888")
    for col_h, h in enumerate(["Severity", "Count", "Financial Impact (₹)"], 1):
        c = ws1.cell(row=4, column=col_h, value=h)
        c.fill = header_fill
        c.font = header_font

    total_all = 0
    total_fin = 0.0
    for i, sev in enumerate(["critical", "high", "medium", "low"], 5):
        count = db.query(func.count(AuditException.id)).filter(
            AuditException.company_id == cid, AuditException.severity == sev,
        ).scalar() or 0
        fin = float(db.query(func.sum(AuditException.financial_impact)).filter(
            AuditException.company_id == cid, AuditException.severity == sev,
        ).scalar() or 0)
        total_all += count
        total_fin += fin

        sev_cell = ws1.cell(row=i, column=1, value=sev.upper())
        color = severity_colors.get(sev, "CCCCCC")
        sev_cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        sev_cell.font = Font(bold=True, color="FFFFFF")
        ws1.cell(row=i, column=2, value=count)
        ws1.cell(row=i, column=3, value=fin)

    ws1.cell(row=9, column=1, value="TOTAL").font = Font(bold=True)
    ws1.cell(row=9, column=2, value=total_all).font = Font(bold=True)
    ws1.cell(row=9, column=3, value=total_fin).font = Font(bold=True, color="C8A951")

    # By business area
    ws1.cell(row=11, column=1, value="BY BUSINESS AREA").font = Font(bold=True, color="888888")
    for col_h, h in enumerate(["Business Area", "Count", "Financial Impact (₹)"], 1):
        c = ws1.cell(row=12, column=col_h, value=h)
        c.fill = header_fill
        c.font = header_font

    area_rows = db.query(
        AuditException.category,
        func.count(AuditException.id),
        func.sum(AuditException.financial_impact),
    ).filter(AuditException.company_id == cid).group_by(AuditException.category).all()

    for j, (cat, cnt, fin_sum) in enumerate(area_rows, 13):
        ws1.cell(row=j, column=1, value=(cat or "").replace("_", " ").title())
        ws1.cell(row=j, column=2, value=cnt)
        ws1.cell(row=j, column=3, value=float(fin_sum or 0))

    for col, w in enumerate([25, 10, 22], 1):
        ws1.column_dimensions[get_column_letter(col)].width = w

    # ── Sheet 2: All Exceptions ───────────────────────────────────────────────
    def _write_exception_sheet(ws, exceptions):
        headers = [
            "Exc #", "Category", "Title", "Severity", "Risk Impact",
            "Financial Impact (₹)", "Status", "Detected On",
            "Description", "Evidence Summary", "AI Analysis",
        ]
        for col, h in enumerate(headers, 1):
            c = ws.cell(row=1, column=col, value=h)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center")

        for i, exc in enumerate(exceptions, 2):
            impact = float(exc.financial_impact or 0)
            ws.cell(row=i, column=1, value=i - 1)
            ws.cell(row=i, column=2, value=exc.category or "")
            ws.cell(row=i, column=3, value=exc.title or "")
            sev_c = ws.cell(row=i, column=4, value=(exc.severity or "").upper())
            color = severity_colors.get((exc.severity or "").lower(), "CCCCCC")
            sev_c.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
            sev_c.font = Font(bold=True, color="FFFFFF", size=9)
            sev_c.alignment = Alignment(horizontal="center")
            ws.cell(row=i, column=5, value=exc.risk_impact or "")
            ws.cell(row=i, column=6, value=impact)
            ws.cell(row=i, column=7, value=exc.status or "")
            ws.cell(row=i, column=8, value=str(exc.detected_on) if exc.detected_on else "")
            ws.cell(row=i, column=9, value=exc.description or "").alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(row=i, column=10, value=_flatten_evidence(exc.supporting_data)).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(row=i, column=11, value=exc.ai_analysis or "").alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[i].height = 55

        col_widths = [6, 18, 45, 11, 20, 22, 16, 14, 50, 45, 50]
        for col, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = w
        ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("All Exceptions")
    all_exc = db.query(AuditException).filter(
        AuditException.company_id == cid
    ).order_by(AuditException.detected_on.desc()).all()
    _write_exception_sheet(ws2, all_exc)

    # ── Sheet 3: Critical & High ──────────────────────────────────────────────
    ws3 = wb.create_sheet("Critical & High")
    critical_exc = db.query(AuditException).filter(
        AuditException.company_id == cid,
        AuditException.severity.in_(["critical", "high"]),
    ).order_by(AuditException.detected_on.desc()).all()
    _write_exception_sheet(ws3, critical_exc)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generate_audit_deck_pdf(db: "Session", company_id: str) -> io.BytesIO:
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.platypus import Table, TableStyle
    from app.models.audit import AuditException
    from app.models.tenant import Company
    from sqlalchemy import func
    import uuid

    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=A4)
    w, h = A4

    company = db.query(Company).filter(Company.id == uuid.UUID(company_id)).first()
    company_name = company.name if company else "Company"
    company_gstin = getattr(company, 'gstin', '') or ''
    cid = uuid.UUID(company_id)

    NAVY = (0.05, 0.07, 0.09)
    GOLD = (0.78, 0.66, 0.32)
    WHITE = (1, 1, 1)
    GRAY = (0.6, 0.6, 0.6)

    def bg(cv):
        cv.setFillColorRGB(*NAVY)
        cv.rect(0, 0, w, h, fill=1, stroke=0)

    # ── Page 1: Cover ─────────────────────────────────────────────────────────
    bg(c)
    c.setFillColorRGB(*GOLD)
    c.setFont("Helvetica-Bold", 26)
    c.drawCentredString(w / 2, h - 120, "INTERNAL AUDIT REPORT")

    c.setFillColorRGB(*WHITE)
    c.setFont("Helvetica-Bold", 16)
    c.drawCentredString(w / 2, h - 160, company_name)

    if company_gstin:
        c.setFillColorRGB(*GRAY)
        c.setFont("Helvetica", 11)
        c.drawCentredString(w / 2, h - 185, f"GSTIN: {company_gstin}")

    c.setFillColorRGB(*GOLD)
    c.setLineWidth(1)
    c.setStrokeColorRGB(*GOLD)
    c.line(60, h - 210, w - 60, h - 210)

    c.setFillColorRGB(*GRAY)
    c.setFont("Helvetica", 10)
    c.drawCentredString(w / 2, h - 230, f"Generated: {__import__('datetime').datetime.now().strftime('%d %B %Y')}")
    c.drawCentredString(w / 2, h - 250, "Prepared by Continuous Audit Engine — Powered by Anthropic Claude")

    c.showPage()

    # ── Page 2: Executive Summary ─────────────────────────────────────────────
    bg(c)
    c.setFillColorRGB(*GOLD)
    c.setFont("Helvetica-Bold", 18)
    c.drawString(50, h - 70, "Executive Summary")
    c.setStrokeColorRGB(*GOLD)
    c.line(50, h - 82, w - 50, h - 82)

    severities = ["critical", "high", "medium", "low"]
    sev_colors_rgb = {
        "critical": (0.9, 0.1, 0.1),
        "high": (0.9, 0.4, 0.0),
        "medium": (0.8, 0.6, 0.0),
        "low": (0.1, 0.6, 0.2),
    }

    y = h - 120
    total_count = 0
    total_impact = 0.0
    sev_data = []
    for sev in severities:
        count = db.query(func.count(AuditException.id)).filter(
            AuditException.company_id == cid, AuditException.severity == sev,
        ).scalar() or 0
        fin = float(db.query(func.sum(AuditException.financial_impact)).filter(
            AuditException.company_id == cid, AuditException.severity == sev,
        ).scalar() or 0)
        total_count += count
        total_impact += fin
        sev_data.append((sev, count, fin))

    c.setFillColorRGB(*WHITE)
    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, y, "Severity")
    c.drawString(200, y, "Exception Count")
    c.drawString(360, y, "Financial Exposure (₹)")
    y -= 8
    c.setStrokeColorRGB(0.3, 0.3, 0.3)
    c.line(50, y, w - 50, y)
    y -= 22

    for sev, count, fin in sev_data:
        rgb = sev_colors_rgb.get(sev, (0.8, 0.8, 0.8))
        c.setFillColorRGB(*rgb)
        c.roundRect(50, y - 4, 90, 18, 3, fill=1, stroke=0)
        c.setFillColorRGB(*WHITE)
        c.setFont("Helvetica-Bold", 10)
        c.drawCentredString(95, y + 3, sev.upper())

        c.setFillColorRGB(*WHITE)
        c.setFont("Helvetica", 11)
        c.drawString(220, y + 2, str(count))
        c.drawString(360, y + 2, f"₹ {fin:,.0f}")
        y -= 28

    y -= 10
    c.setStrokeColorRGB(0.3, 0.3, 0.3)
    c.line(50, y, w - 50, y)
    y -= 22
    c.setFillColorRGB(*GOLD)
    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, y, "TOTAL")
    c.drawString(220, y, str(total_count))
    c.drawString(360, y, f"₹ {total_impact:,.0f}")

    c.showPage()

    # ── Page 3+: Top Findings ──────────────────────────────────────────────────
    top_exceptions = (
        db.query(AuditException)
        .filter(AuditException.company_id == cid)
        .order_by(AuditException.financial_impact.desc().nullslast(), AuditException.detected_on.desc())
        .limit(20)
        .all()
    )

    bg(c)
    c.setFillColorRGB(*GOLD)
    c.setFont("Helvetica-Bold", 18)
    c.drawString(50, h - 70, "Top Audit Findings")
    c.setStrokeColorRGB(*GOLD)
    c.line(50, h - 82, w - 50, h - 82)

    y = h - 115
    line_h = 14

    for idx, exc in enumerate(top_exceptions, 1):
        if y < 80:
            c.showPage()
            bg(c)
            c.setFillColorRGB(*GOLD)
            c.setFont("Helvetica-Bold", 14)
            c.drawString(50, h - 50, "Top Audit Findings (continued)")
            c.line(50, h - 62, w - 50, h - 62)
            y = h - 90

        # Finding header line
        sev = (exc.severity or "low").lower()
        rgb = sev_colors_rgb.get(sev, (0.5, 0.5, 0.5))
        c.setFillColorRGB(*rgb)
        c.roundRect(50, y - 3, 60, 14, 2, fill=1, stroke=0)
        c.setFillColorRGB(*WHITE)
        c.setFont("Helvetica-Bold", 8)
        c.drawCentredString(80, y + 1, sev.upper())

        c.setFillColorRGB(*WHITE)
        c.setFont("Helvetica-Bold", 10)
        title = (exc.title or "")[:80]
        c.drawString(120, y + 1, f"{idx}. {title}")

        fin = float(exc.financial_impact or 0)
        if fin:
            c.setFillColorRGB(*GOLD)
            c.setFont("Helvetica", 9)
            c.drawRightString(w - 50, y + 1, f"₹ {fin:,.0f}")

        y -= line_h

        # Category + date
        area = (exc.category or "").replace("_", " ").title()
        date_str = str(exc.detected_on) if exc.detected_on else ""
        c.setFillColorRGB(*GRAY)
        c.setFont("Helvetica", 8)
        c.drawString(120, y + 1, f"{area}  |  Detected: {date_str}  |  Status: {exc.status or ''}")
        y -= line_h

        # Description
        if exc.description:
            desc = exc.description[:200]
            c.setFillColorRGB(0.8, 0.8, 0.8)
            c.setFont("Helvetica", 8)
            # Wrap at ~100 chars
            words = desc.split()
            line_buf = []
            for word in words:
                if sum(len(w) + 1 for w in line_buf) + len(word) > 100:
                    c.drawString(120, y + 1, " ".join(line_buf))
                    y -= 11
                    if y < 80:
                        break
                    line_buf = [word]
                else:
                    line_buf.append(word)
            if line_buf:
                c.drawString(120, y + 1, " ".join(line_buf))
                y -= 11

        # Divider
        c.setStrokeColorRGB(0.2, 0.2, 0.2)
        c.line(50, y, w - 50, y)
        y -= 10

    c.showPage()
    c.save()
    buffer.seek(0)
    return buffer
