"""Engagement reports: the Statutory audit working paper (CARO 2020 / tax audit / IFC matrices)
and the Internal Audit Report (process-wise observations with management action plan).

Everything here is staged for the auditor: a matrix row with no CAE finding says so — it is
not a conclusion that the clause is clean.
"""
from __future__ import annotations

import io
import re
from collections import defaultdict
from datetime import datetime

from sqlalchemy.orm import Session

from app.statutory import reporting

# where a clause's reportable amount is not the findings' impact column
_CLAUSE_NOTES = {
    ("3CD", "26"): "impact column = s.16 interest; the s.43B(h) amount for the year is on Statutory > MSME interest",
    ("3CD", "22"): "impact column = s.16 interest inadmissible u/s 23 MSMED",
}
IA_RATING = {"critical": "High", "high": "High", "medium": "Medium", "low": "Low", "info": "Low"}
_SEV = ("critical", "high", "medium", "low")


def inr(x, symbol: str = "₹") -> str:
    """Indian digit grouping: 12,34,56,789."""
    if x is None:
        return "—"
    n = round(float(x))
    sign, digits = ("-" if n < 0 else ""), str(abs(n))
    if len(digits) > 3:
        head, tail = digits[:-3], digits[-3:]
        head = ",".join(re.findall(r"\d{1,2}", head[::-1]))[::-1]
        digits = f"{head},{tail}"
    return f"{sign}{symbol}{digits}"


def _rows(db: Session, engagement):
    from app.models.audit import AuditException, AuditRule

    q = (db.query(AuditException, AuditRule.rule_code, AuditRule.rule_name, AuditRule.remediation_guidance)
         .outerjoin(AuditRule, AuditRule.id == AuditException.rule_id)
         .filter(AuditException.engagement_id == engagement.id))
    rank = {s: i for i, s in enumerate(_SEV)}
    return sorted(q.all(), key=lambda r: (rank.get(r[0].severity, 9), -float(r[0].financial_impact or 0)))


def _amount(exc) -> float:
    return float(exc.financial_impact or 0)


def build_summary(db: Session, engagement) -> dict:
    from app.models.audit import AuditRule

    rows = _rows(db, engagement)
    materiality = float(engagement.materiality) if engagement.materiality else None
    by_sev = {s: 0 for s in _SEV}
    for exc, *_ in rows:
        by_sev[exc.severity] = by_sev.get(exc.severity, 0) + 1
    out = {
        "engagement": {"id": str(engagement.id), "label": engagement.label, "audit_type": engagement.audit_type,
                       "period_from": engagement.period_from.isoformat(), "period_to": engagement.period_to.isoformat(),
                       "financial_year": engagement.financial_year, "materiality": materiality,
                       "caro_applicable": engagement.caro_applicable, "ifc_applicable": engagement.ifc_applicable,
                       "tax_audit_applicable": engagement.tax_audit_applicable, "status": engagement.status},
        "total_findings": len(rows), "by_severity": by_sev,
        "total_financial_impact": round(sum(_amount(e) for e, *_ in rows), 2),
        "material_findings": sum(1 for e, *_ in rows if materiality and _amount(e) >= materiality),
    }
    rules = db.query(AuditRule).filter(AuditRule.company_id == engagement.company_id).order_by(AuditRule.rule_code).all()
    counts = defaultdict(int)
    for _, code, *_ in rows:
        counts[code] += 1
    out["scope"] = [{"rule_code": r.rule_code, "rule_name": r.rule_name, "active": r.is_active,
                     "in_scope": r.is_active and reporting.in_scope(r.rule_code, engagement),
                     "findings": counts.get(r.rule_code, 0)} for r in rules]

    if engagement.audit_type == "statutory":
        matrix = []
        frames = [("CARO 2020", "CARO", reporting.CARO_2020, engagement.caro_applicable),
                  (reporting.tax_audit_frame(engagement.financial_year), "3CD", reporting.FORM_3CD,
                   engagement.tax_audit_applicable)]
        for label, key, clauses, applies in frames:
            if not applies:
                continue
            for clause, title in clauses.items():
                hits = [(e, code) for e, code, *_ in rows
                        if any(f == key and ref == clause for f, ref, _ in reporting.RULE_MAP.get(code, {}).get("refs", []))]
                matrix.append({
                    "frame": label, "clause": clause if key == "CARO" else f"cl. {clause}", "title": title,
                    "findings": len(hits), "financial_impact": round(sum(_amount(e) for e, _ in hits), 2),
                    "rules": sorted({c for _, c in hits}),
                    "draft_position": ("Exceptions found — evaluate and report" if hits
                                       else "No exception from CAE tests — auditor to conclude on other evidence")
                                      + (f" ({_CLAUSE_NOTES[(key, clause)]})" if (key, clause) in _CLAUSE_NOTES else ""),
                })
        out["reporting_matrix"] = matrix
        if engagement.ifc_applicable:
            ifc = defaultdict(lambda: {"findings": 0, "financial_impact": 0.0, "rule_name": ""})
            for e, code, name, _ in rows:
                if any(f == "IFC" for f, *_ in reporting.RULE_MAP.get(code, {}).get("refs", [])):
                    ifc[code]["findings"] += 1
                    ifc[code]["financial_impact"] += _amount(e)
                    ifc[code]["rule_name"] = name
            out["ifc_deficiencies"] = [{"rule_code": k, **v, "classification": "To be assessed (deficiency / significant / material weakness)"}
                                       for k, v in sorted(ifc.items())]
    else:
        procs = {p: {"process": p, "findings": 0, "high": 0, "medium": 0, "low": 0, "financial_impact": 0.0, "open": 0}
                 for p in reporting.PROCESSES}
        for e, code, *_ in rows:
            p = procs.setdefault(reporting.process_for(code), {"process": reporting.process_for(code), "findings": 0,
                                                               "high": 0, "medium": 0, "low": 0, "financial_impact": 0.0, "open": 0})
            p["findings"] += 1
            p[IA_RATING.get(e.severity, "Low").lower()] += 1
            p["financial_impact"] += _amount(e)
            p["open"] += e.status in ("open", "in_review", "management_query_sent", "remediation_in_progress")
        out["by_process"] = [dict(v, financial_impact=round(v["financial_impact"], 2)) for v in procs.values()]
    return out


# ---------------------------------------------------------------------------- Excel helpers
def _wb():
    import openpyxl
    return openpyxl.Workbook()


def _style_header(ws, row: int, ncols: int):
    from openpyxl.styles import Alignment, Font, PatternFill
    fill = PatternFill(start_color="0D1117", end_color="0D1117", fill_type="solid")
    for col in range(1, ncols + 1):
        c = ws.cell(row=row, column=col)
        c.fill = fill
        c.font = Font(bold=True, color="C8A951", size=10)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _table(ws, start_row: int, headers: list[str], rows: list[list], widths: list[int] | None = None):
    from openpyxl.styles import Alignment
    for i, h in enumerate(headers, 1):
        ws.cell(row=start_row, column=i, value=h)
    _style_header(ws, start_row, len(headers))
    for r, values in enumerate(rows, start_row + 1):
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(v, float):
                cell.number_format = "#,##0.00"
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


def _cover(ws, title: str, engagement, company, extra: list[tuple[str, object]]):
    from openpyxl.styles import Font
    ws.title = "Cover"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 70
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=16, color="C8A951")
    rows = [("Entity", company.name if company else ""), ("GSTIN", getattr(company, "gstin", "") or ""),
            ("CIN", getattr(company, "cin", "") or ""), ("Engagement", engagement.label),
            ("Period", f"{engagement.period_from:%d-%m-%Y} to {engagement.period_to:%d-%m-%Y}"),
            ("Status", engagement.status), *extra,
            ("Generated", datetime.now().strftime("%d %b %Y %H:%M")),
            ("Note", "Prepared from CAE automated tests. Every finding and matrix row is staged for the "
                     "auditor's review — CAE does not conclude or sign.")]
    for i, (k, v) in enumerate(rows, 3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True, color="666666")
        ws.cell(row=i, column=2, value=v)


def _refs_text(code, engagement) -> str:
    return "; ".join(f"{r['frame']} {r['ref']}" + (f" ({r['note']})" if r["note"] else "")
                     for r in reporting.refs_for(code, engagement))


def _evidence(data) -> str:
    if not data:
        return ""
    if isinstance(data, dict):
        return " | ".join(f"{k}: {v}" for k, v in list(data.items())[:10])
    return str(data)[:300]


def statutory_working_paper(db: Session, engagement) -> io.BytesIO:
    from app.models.tenant import Company

    company = db.query(Company).filter(Company.id == engagement.company_id).first()
    summary = build_summary(db, engagement)
    rows = _rows(db, engagement)
    materiality = float(engagement.materiality) if engagement.materiality else None
    yes = lambda b: "Applicable" if b else "Not applicable"  # noqa: E731

    wb = _wb()
    _cover(wb.active, "STATUTORY AUDIT — WORKING PAPER", engagement, company, [
        ("Reporting framework", "Companies Act 2013 s.143; Standards on Auditing"),
        ("CARO 2020", yes(engagement.caro_applicable)),
        ("IFC reporting s.143(3)(i)", yes(engagement.ifc_applicable)),
        ("Tax audit", f"{yes(engagement.tax_audit_applicable)} — {reporting.tax_audit_frame(engagement.financial_year)}"),
        ("Overall materiality (SA 320)", inr(materiality) if materiality else "Not set"),
        ("Findings", f"{summary['total_findings']} (critical {summary['by_severity'].get('critical', 0)}, "
                     f"high {summary['by_severity'].get('high', 0)}); at or above materiality: {summary['material_findings']}"),
        ("Total financial impact", inr(summary['total_financial_impact'])),
    ])

    ws = wb.create_sheet("Findings")
    _table(ws, 1, ["#", "Rule", "Finding", "Severity", "Impact (₹)", "≥ Materiality", "Reporting reference",
                   "Status", "Description", "Evidence"],
           [[i, code, e.title, e.severity.upper(), _amount(e),
             ("Yes" if materiality and _amount(e) >= materiality else "No") if materiality else "—",
             _refs_text(code, engagement), e.status, e.description, _evidence(e.supporting_data)]
            for i, (e, code, *_) in enumerate(rows, 1)],
           [5, 9, 46, 10, 15, 12, 46, 14, 60, 60])

    matrix = summary.get("reporting_matrix", [])
    for frame in dict.fromkeys(m["frame"] for m in matrix):
        ws = wb.create_sheet("CARO 2020" if frame == "CARO 2020" else "Tax audit")
        _table(ws, 1, ["Clause", "Requirement", "CAE findings", "Impact (₹)", "Rules", "Draft position"],
               [[m["clause"], m["title"], m["findings"], m["financial_impact"], ", ".join(m["rules"]), m["draft_position"]]
                for m in matrix if m["frame"] == frame], [12, 60, 12, 16, 22, 50])
    if summary.get("ifc_deficiencies") is not None:
        ws = wb.create_sheet("IFC")
        _table(ws, 1, ["Rule", "Control tested", "Exceptions", "Impact (₹)", "Classification"],
               [[d["rule_code"], d["rule_name"], d["findings"], round(d["financial_impact"], 2), d["classification"]]
                for d in summary["ifc_deficiencies"]], [10, 46, 12, 16, 50])
    ws = wb.create_sheet("Scope")
    _table(ws, 1, ["Rule", "Test", "Active", "In statutory scope", "Findings"],
           [[s["rule_code"], s["rule_name"], "Yes" if s["active"] else "No", "Yes" if s["in_scope"] else "No", s["findings"]]
            for s in summary["scope"]], [10, 50, 8, 18, 10])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _workflow(db: Session, exception_ids):
    from app.models.workflow import ManagementQuery, QueryResponse, RemediationAction

    responses, actions = {}, {}
    if not exception_ids:
        return responses, actions
    for q, r in (db.query(ManagementQuery, QueryResponse).outerjoin(QueryResponse, QueryResponse.query_id == ManagementQuery.id)
                 .filter(ManagementQuery.exception_id.in_(exception_ids)).order_by(QueryResponse.responded_at).all()):
        if r is not None:
            responses[q.exception_id] = f"{r.response_text} — {r.responded_by_name}, {r.responded_at:%d-%m-%Y}"
        else:
            responses.setdefault(q.exception_id, f"Query {q.query_no} raised; response due {q.due_date:%d-%m-%Y}")
    for a in db.query(RemediationAction).filter(RemediationAction.exception_id.in_(exception_ids)).all():
        actions[a.exception_id] = a
    return responses, actions


def internal_audit_report(db: Session, engagement) -> io.BytesIO:
    from app.models.tenant import Company

    company = db.query(Company).filter(Company.id == engagement.company_id).first()
    summary = build_summary(db, engagement)
    rows = _rows(db, engagement)
    responses, actions = _workflow(db, [e.id for e, *_ in rows])

    wb = _wb()
    _cover(wb.active, "INTERNAL AUDIT REPORT", engagement, company, [
        ("Basis", "ICAI Standards on Internal Audit — observations rated High / Medium / Low"),
        ("Observations", f"{summary['total_findings']}"),
        ("Total financial impact", inr(summary['total_financial_impact'])),
    ])
    ws = wb.create_sheet("Summary by process")
    _table(ws, 1, ["Process", "Observations", "High", "Medium", "Low", "Open", "Impact (₹)"],
           [[p["process"], p["findings"], p["high"], p["medium"], p["low"], p["open"], p["financial_impact"]]
            for p in summary["by_process"]], [30, 14, 8, 10, 8, 8, 18])
    ws = wb.create_sheet("Observations")
    obs = []
    for i, (e, code, name, guidance) in enumerate(rows, 1):
        act = actions.get(e.id)
        obs.append([f"IA-{i:03d}", reporting.process_for(code), e.title, IA_RATING.get(e.severity, "Low"), _amount(e),
                    e.description, e.risk_impact or "", guidance or "", responses.get(e.id, "Awaited"),
                    act.responsible_name if act else "", act.target_date.isoformat() if act and act.target_date else "",
                    (act.status if act else e.status), code])
    _table(ws, 1, ["Obs #", "Process", "Observation", "Risk rating", "Impact (₹)", "Details", "Risk / implication",
                   "Recommendation", "Management response", "Action owner", "Target date", "Status", "Test"],
           obs, [9, 22, 44, 11, 15, 60, 18, 44, 44, 18, 12, 14, 9])
    ws = wb.create_sheet("Action plan")
    plan = [[f"IA-{i:03d}", e.title, a.action_plan, a.responsible_name, a.target_date.isoformat() if a.target_date else "",
             a.status, a.retest_outcome or ""]
            for i, (e, *_r) in enumerate(rows, 1) if (a := actions.get(e.id))]
    _table(ws, 1, ["Obs #", "Observation", "Agreed action", "Owner", "Target date", "Status", "Re-test"], plan,
           [9, 44, 50, 20, 12, 12, 14])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------------------- PDF
def engagement_pdf(db: Session, engagement) -> io.BytesIO:
    """Short PDF: cover, summary, the statutory reporting matrix or the process summary, top findings."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from app.models.tenant import Company

    company = db.query(Company).filter(Company.id == engagement.company_id).first()
    summary = build_summary(db, engagement)
    rows = _rows(db, engagement)
    statutory = engagement.audit_type == "statutory"
    styles = getSampleStyleSheet()
    small = styles["BodyText"].clone("small", fontSize=8, leading=10)
    title = "Statutory Audit — Findings and Reporting Matrix" if statutory else "Internal Audit Report"

    def tbl(data, widths):
        t = Table([[Paragraph(str(c), small) for c in r] for r in data], colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0D1117")),
                               ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#C8A951")),
                               ("GRID", (0, 0), (-1, -1), 0.25, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        return t

    story = [Paragraph(title, styles["Title"]),
             Paragraph(f"{company.name if company else ''} — {engagement.label}", styles["Heading2"]),
             Paragraph(f"Period {engagement.period_from:%d-%m-%Y} to {engagement.period_to:%d-%m-%Y} · generated "
                       f"{datetime.now():%d %b %Y}", styles["BodyText"]), Spacer(1, 0.4 * cm)]
    sev = summary["by_severity"]
    facts = [["Findings", "Critical", "High", "Medium", "Low", "Impact"],
             [summary["total_findings"], sev.get("critical", 0), sev.get("high", 0), sev.get("medium", 0),
              sev.get("low", 0), inr(summary['total_financial_impact'], "Rs.")]]
    story += [tbl(facts, [2.6 * cm] * 5 + [3.6 * cm]), Spacer(1, 0.4 * cm)]
    if statutory:
        mat = summary.get("materiality")
        story.append(Paragraph(
            f"CARO 2020: {'applicable' if engagement.caro_applicable else 'not applicable'} · IFC: "
            f"{'applicable' if engagement.ifc_applicable else 'not applicable'} · Tax audit: "
            f"{'applicable' if engagement.tax_audit_applicable else 'not applicable'} · Materiality: "
            f"{inr(mat, 'Rs.') if mat else 'not set'} ({summary['material_findings']} findings at or above)", small))
        if summary.get("reporting_matrix"):
            story += [Spacer(1, 0.3 * cm), Paragraph("Reporting matrix", styles["Heading3"]),
                      tbl([["Frame", "Clause", "Requirement", "Findings", "Impact", "Draft position"]] +
                          [[m["frame"].split(" (")[0], m["clause"], m["title"], m["findings"], inr(m['financial_impact'], "Rs."),
                            m["draft_position"]] for m in summary["reporting_matrix"]],
                          [2.4 * cm, 1.6 * cm, 5.4 * cm, 1.4 * cm, 2.2 * cm, 4.6 * cm])]
    else:
        story += [Paragraph("Observations by process", styles["Heading3"]),
                  tbl([["Process", "Obs.", "High", "Medium", "Low", "Open", "Impact"]] +
                      [[p["process"], p["findings"], p["high"], p["medium"], p["low"], p["open"], inr(p['financial_impact'], "Rs.")]
                       for p in summary["by_process"]], [5 * cm, 1.5 * cm, 1.5 * cm, 1.7 * cm, 1.5 * cm, 1.5 * cm, 3 * cm])]
    story += [PageBreak(), Paragraph("Key findings" if statutory else "Key observations", styles["Heading2"])]
    top = [["#", "Finding", "Rating", "Impact", "Reference" if statutory else "Process"]]
    for i, (e, code, *_) in enumerate(rows[:40], 1):
        top.append([i, e.title, e.severity.upper() if statutory else IA_RATING.get(e.severity, "Low"),
                    inr(_amount(e), "Rs."), _refs_text(code, engagement) if statutory else reporting.process_for(code)])
    story += [tbl(top, [0.8 * cm, 7 * cm, 1.8 * cm, 2.4 * cm, 5.6 * cm]), Spacer(1, 0.4 * cm),
              Paragraph("Staged for the auditor's review — CAE flags exceptions; it does not conclude or sign.", small)]
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.5 * cm,
                      bottomMargin=1.5 * cm, title=title).build(story)
    buf.seek(0)
    return buf
