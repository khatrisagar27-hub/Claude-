"""
Build a premium Excel workbook for Foreign Exchange Gain/(Loss) calculation
on Sales and Purchase transactions.

Design goals
------------
* Input sheets are the base: Invoices and Receipts/Payments are typed by the user.
* An Allocation sheet resolves the many-to-many reality of forex settlement:
    - one invoice settled by 3+ receipts/payments across different dates
    - part advances (receipt/payment dated before the invoice)
    - one receipt/payment tied to 2-3 invoices
* Everything downstream is driven by live Excel formulas + structured table
  references, so the workbook recalculates as the user edits inputs.
* Realised gain/(loss) is computed per allocation; unrealised (revaluation of
  the open foreign-currency balance at the closing rate) is computed per invoice.

Author: generated for the Foreign Gain/(Loss) workbook task.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
import datetime

# --------------------------------------------------------------------------- #
#  Palette / shared styles
# --------------------------------------------------------------------------- #
NAVY      = "1F3864"   # titles
BLUE      = "2E5496"   # input headers
SLATE     = "44546A"   # computed headers
TEAL      = "0F6E6E"   # settlement headers
GOLD      = "BF8F00"   # accents
LIGHT     = "DDEBF7"   # input cell fill
GREY      = "EDEDED"   # computed cell fill
BAND      = "F2F2F2"
WHITE     = "FFFFFF"
GREEN_F   = "548235"   # gain
RED_F     = "C00000"   # loss
GREEN_BG  = "E2EFDA"
RED_BG    = "FCE4E4"

FMT_AMT   = '#,##0.00'
FMT_FC    = '#,##0.00'
FMT_RATE  = '0.0000'
FMT_DATE  = 'dd-mmm-yyyy'
FMT_INT   = '#,##0'

thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def title_font(sz=16, color=WHITE):
    return Font(name="Calibri", size=sz, bold=True, color=color)

def hdr_font(color=WHITE):
    return Font(name="Calibri", size=10, bold=True, color=color)

def cell_font(bold=False, color="000000", italic=False):
    return Font(name="Calibri", size=10, bold=bold, color=color, italic=italic)

def fill(color):
    return PatternFill("solid", fgColor=color)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT   = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT  = Alignment(horizontal="right", vertical="center")
LEFT_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True)

wb = Workbook()

# --------------------------------------------------------------------------- #
#  Helper to draw a banner title row
# --------------------------------------------------------------------------- #
def banner(ws, text, subtitle, ncols, color=NAVY):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=text)
    c.font = title_font(16)
    c.fill = fill(color)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    s = ws.cell(row=2, column=2 if False else 1, value=subtitle)
    s.font = Font(name="Calibri", size=9, italic=True, color="FFFFFF")
    s.fill = fill(color)
    s.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 16


def write_table(ws, top_row, columns, data, table_name, style="TableStyleMedium2",
                header_fill=BLUE, input_cols=None, computed_cols=None):
    """columns: list of header strings. data: list of rows (lists). Returns table ref."""
    ncols = len(columns)
    # header
    for j, h in enumerate(columns, start=1):
        c = ws.cell(row=top_row, column=j, value=h)
        c.font = hdr_font()
        c.fill = fill(header_fill)
        c.alignment = CENTER
        c.border = BORDER
    ws.row_dimensions[top_row].height = 30
    # data
    for i, row in enumerate(data, start=top_row + 1):
        for j, val in enumerate(row, start=1):
            c = ws.cell(row=i, column=j, value=val)
            c.border = BORDER
            c.alignment = LEFT if isinstance(val, str) and not str(val).startswith("=") else RIGHT
    last_row = top_row + len(data)
    ref = f"A{top_row}:{get_column_letter(ncols)}{last_row}"
    tbl = Table(displayName=table_name, ref=ref)
    tbl.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True,
                                        showColumnStripes=False,
                                        showFirstColumn=False, showLastColumn=False)
    ws.add_table(tbl)
    return last_row

# =========================================================================== #
#  1. READ ME
# =========================================================================== #
ws = wb.active
ws.title = "Read Me"
ws.sheet_view.showGridLines = False
banner(ws, "FOREIGN EXCHANGE GAIN / (LOSS) WORKBOOK",
        "Realised & unrealised forex on sales and purchase transactions  •  driven by live formulas", 8, NAVY)

readme_blocks = [
    ("HOW THIS WORKBOOK IS ORGANISED", NAVY, [
        "1.  Settings  –  set your base (reporting) currency, the reporting / closing date and the closing exchange rate per foreign currency.",
        "2.  Invoices  –  enter every SALES and PURCHASE invoice raised/received in foreign currency (the base input).",
        "3.  Receipts & Payments  –  enter every money movement: receipts against sales, payments against purchases (incl. advances).",
        "4.  Allocation & Forex Calc  –  link each receipt/payment to the invoice(s) it settles. This is where gain/(loss) is calculated.",
        "5.  Dashboard  –  automatic summary of realised, unrealised and net forex impact by transaction type and party.",
    ]),
    ("WHY THE ALLOCATION SHEET EXISTS (read this)", GOLD, [
        "Forex settlement is many-to-many. The Allocation sheet is the bridge that captures every real-world situation:",
        "•  One invoice settled by 3 or more receipts/payments on different dates  →  enter one allocation line per date.",
        "•  A part advance (money received/paid BEFORE the invoice date)  →  the line is auto-tagged 'Advance'.",
        "•  One single receipt/payment settling 2–3 invoices  →  enter one allocation line per invoice, splitting the amount.",
        "Each allocation line earns its own gain/(loss), because each settlement happened at its own exchange rate.",
    ]),
    ("HOW GAIN / (LOSS) IS CALCULATED", TEAL, [
        "Realised (per allocation line, when money actually moves):",
        "      Forex Gain/(Loss) = FC Allocated × (Settlement Rate − Invoice Rate) × sign",
        "      sign = +1 for SALES   |   sign = −1 for PURCHASE",
        "      → Sales: you gain when the currency strengthens after invoicing; Purchase: you gain when it weakens.",
        "Unrealised (per invoice, on the still-open foreign-currency balance at the closing rate):",
        "      Unrealised Gain/(Loss) = FC Outstanding × (Closing Rate − Invoice Rate) × sign",
        "A POSITIVE figure is a GAIN; a figure in (brackets)/red is a LOSS to the company.",
    ]),
    ("CONVENTIONS", SLATE, [
        "•  Rate = units of base currency per 1 unit of foreign currency (e.g. INR per 1 USD).",
        "•  Blue headers = your input.   Grey headers = automatically calculated, do not overtype.",
        "•  Tables auto-extend: start typing on the row below a table and the formatting & formulas copy down.",
        "•  Dropdowns are provided for Type and Currency; Invoice No / Settlement Ref pick-lists are validated.",
        "•  The sample rows are a worked example covering every scenario above – clear them to start your own data.",
    ]),
]
r = 4
for heading, color, lines in readme_blocks:
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    c = ws.cell(row=r, column=1, value=heading)
    c.font = Font(bold=True, size=11, color=WHITE)
    c.fill = fill(color)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[r].height = 20
    r += 1
    for line in lines:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        c = ws.cell(row=r, column=1, value=line)
        c.font = cell_font()
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
        ws.row_dimensions[r].height = 15 if len(line) < 95 else 28
        r += 1
    r += 1
ws.column_dimensions["A"].width = 14
for col in "BCDEFGH":
    ws.column_dimensions[col].width = 14

# =========================================================================== #
#  2. SETTINGS
# =========================================================================== #
ws = wb.create_sheet("Settings")
ws.sheet_view.showGridLines = False
banner(ws, "SETTINGS  &  CURRENCY MASTER",
        "Set once at period start, then revisit only the closing rates at period end", 6, SLATE)

ws.cell(row=4, column=1, value="Company / Entity Name").font = cell_font(bold=True)
ws.cell(row=4, column=3, value="Your Company Pvt Ltd")
ws.cell(row=5, column=1, value="Base / Reporting Currency").font = cell_font(bold=True)
ws.cell(row=5, column=3, value="INR")
ws.cell(row=6, column=1, value="Reporting / Closing Date").font = cell_font(bold=True)
ws.cell(row=6, column=3, value=datetime.date(2026, 3, 31)).number_format = FMT_DATE
for rr in range(4, 7):
    lab = ws.cell(row=rr, column=1)
    lab.fill = fill(LIGHT)
    val = ws.cell(row=rr, column=3)
    val.fill = fill(WHITE); val.border = BORDER; val.font = cell_font(bold=True, color=BLUE)
    ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=2)
    ws.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=4)

# Currency master with closing rates
cur_cols = ["Currency", "Description", "Closing Rate"]
cur_data = [
    ["USD", "US Dollar",        83.50],
    ["EUR", "Euro",             90.20],
    ["GBP", "Pound Sterling",  105.00],
    ["AED", "UAE Dirham",       22.70],
    ["JPY", "Japanese Yen",      0.5600],
]
ws.cell(row=8, column=1, value="CLOSING EXCHANGE RATES (base currency per 1 unit of foreign currency)").font = Font(bold=True, color=SLATE, size=10)
write_table(ws, 9, cur_cols, cur_data, "tblCurrency", style="TableStyleMedium4", header_fill=TEAL)
for rr in range(10, 10 + len(cur_data)):
    ws.cell(row=rr, column=3).number_format = FMT_RATE
ws.column_dimensions["A"].width = 24
ws.column_dimensions["B"].width = 22
ws.column_dimensions["C"].width = 16
ws.column_dimensions["D"].width = 12

# Defined names
wb.defined_names.add(DefinedName("BaseCurrency", attr_text="Settings!$C$5"))
wb.defined_names.add(DefinedName("ReportingDate", attr_text="Settings!$C$6"))
wb.defined_names.add(DefinedName("CompanyName", attr_text="Settings!$C$4"))

# =========================================================================== #
#  3. INVOICES  (input + per-invoice analytics)
# =========================================================================== #
ws = wb.create_sheet("Invoices")
ws.sheet_view.showGridLines = False
banner(ws, "SALES & PURCHASE INVOICES",
        "Blue = your input    |    Grey = auto-calculated settlement & forex position", 16, BLUE)

inv_columns = [
    "Invoice No", "Transaction Type", "Party Name", "Currency", "Invoice Date",
    "FC Amount", "Invoice Rate",                       # input (A-G)
    "Invoice Value (Base)", "FC Settled", "FC Outstanding", "Settlement Status",
    "Closing Rate", "Realised Gain/(Loss)", "Unrealised Gain/(Loss)",
    "Total Forex Impact", "Remarks / Additional details",
]

# Sample invoices
inv_data = [
    ["INV-S-001", "Sales",    "Alpha Corp (USA)",     "USD", datetime.date(2026,1,10), 30000, 82.00],
    ["INV-S-002", "Sales",    "Beta Ltd (USA)",       "USD", datetime.date(2026,2,5),  12000, 82.50],
    ["INV-S-003", "Sales",    "Beta Ltd (USA)",       "USD", datetime.date(2026,2,20),  8000, 83.10],
    ["INV-S-004", "Sales",    "Gamma Trading (UK)",   "GBP", datetime.date(2026,3,5),   5000, 104.00],
    ["INV-P-001", "Purchase", "Euro Supplies GmbH",   "EUR", datetime.date(2026,1,15), 20000, 89.00],
    ["INV-P-002", "Purchase", "Euro Supplies GmbH",   "EUR", datetime.date(2026,2,25), 10000, 89.80],
]

# build full rows with formulas (one row per invoice)
inv_rows = []
for k, base in enumerate(inv_data):
    inv_no, ttype, party, cur, idate, fc, rate = base
    inv_rows.append(base + [
        "=[@[FC Amount]]*[@[Invoice Rate]]",
        "=SUMIFS(tblAllocations[FC Allocated],tblAllocations[Invoice No],[@[Invoice No]])",
        "=[@[FC Amount]]-[@[FC Settled]]",
        '=IF(ROUND([@[FC Outstanding]],2)=0,"Fully Settled",IF([@[FC Settled]]=0,"Open",IF([@[FC Outstanding]]<0,"Over-allocated","Partly Settled")))',
        "=IFERROR(INDEX(tblCurrency[Closing Rate],MATCH([@[Currency]],tblCurrency[Currency],0)),\"\")",
        "=SUMIFS(tblAllocations[Forex Gain/(Loss)],tblAllocations[Invoice No],[@[Invoice No]])",
        '=IF([@[Closing Rate]]="","",[@[FC Outstanding]]*([@[Closing Rate]]-[@[Invoice Rate]])*IF([@[Transaction Type]]="Sales",1,-1))',
        '=[@[Realised Gain/(Loss)]]+IF(ISNUMBER([@[Unrealised Gain/(Loss)]]),[@[Unrealised Gain/(Loss)]],0)',
        "",
    ])

last = write_table(ws, 4, inv_columns, inv_rows, "tblInvoices",
                   style="TableStyleMedium2", header_fill=BLUE)

# recolor computed headers grey, settlement-amount formatting etc.
computed_idx = list(range(8, 17))   # columns H..P (1-based) -> indices 8..16, remarks 16 input
for j in range(8, 16):              # H..O computed
    ws.cell(row=4, column=j).fill = fill(SLATE)
ws.cell(row=4, column=16).fill = fill(GOLD)  # remarks input accent

# number formats
def fmt_col(ws, col_letter, fmt, r0, r1):
    for rr in range(r0, r1 + 1):
        ws[f"{col_letter}{rr}"].number_format = fmt

r0, r1 = 5, last
fmt_col(ws, "E", FMT_DATE, r0, r1)         # invoice date
fmt_col(ws, "F", FMT_FC,   r0, r1)         # FC amount
fmt_col(ws, "G", FMT_RATE, r0, r1)         # invoice rate
fmt_col(ws, "H", FMT_AMT,  r0, r1)         # base value
fmt_col(ws, "I", FMT_FC,   r0, r1)         # fc settled
fmt_col(ws, "J", FMT_FC,   r0, r1)         # fc outstanding
fmt_col(ws, "L", FMT_RATE, r0, r1)         # closing rate
for col in ["M", "N", "O"]:
    fmt_col(ws, col, FMT_AMT, r0, r1)

# widths
inv_widths = [12,15,22,9,13,13,11,16,12,13,14,11,16,18,15,26]
for j, w in enumerate(inv_widths, start=1):
    ws.column_dimensions[get_column_letter(j)].width = w
ws.freeze_panes = "B5"

# conditional formatting on M, N, O (gain green / loss red)
for col in ["M", "N", "O"]:
    rng = f"{col}{r0}:{col}{r1}"
    ws.conditional_formatting.add(rng,
        CellIsRule(operator="greaterThan", formula=["0"], font=Font(color=GREEN_F, bold=True)))
    ws.conditional_formatting.add(rng,
        CellIsRule(operator="lessThan", formula=["0"], font=Font(color=RED_F, bold=True)))

# =========================================================================== #
#  4. RECEIPTS & PAYMENTS  (settlements input)
# =========================================================================== #
ws = wb.create_sheet("Receipts & Payments")
ws.sheet_view.showGridLines = False
banner(ws, "RECEIPTS & PAYMENTS  (Settlements)",
        "Every money movement in foreign currency — receipts settle sales, payments settle purchases, incl. advances", 12, TEAL)

stl_columns = [
    "Settlement Ref", "Settlement Type", "Party Name", "Currency", "Settlement Date",
    "FC Amount", "Settlement Rate",                         # input A-G
    "Settlement Value (Base)", "FC Allocated", "FC Unallocated", "Allocation Status",
    "Narration",
]
stl_data = [
    ["REC-001", "Receipt", "Alpha Corp (USA)",   "USD", datetime.date(2025,12,20), 10000, 81.50],
    ["REC-002", "Receipt", "Alpha Corp (USA)",   "USD", datetime.date(2026,2,15),  12000, 83.00],
    ["REC-003", "Receipt", "Alpha Corp (USA)",   "USD", datetime.date(2026,3,20),   8000, 83.40],
    ["REC-004", "Receipt", "Beta Ltd (USA)",     "USD", datetime.date(2026,3,10),  20000, 83.00],
    ["REC-005", "Receipt", "Gamma Trading (UK)", "GBP", datetime.date(2026,3,25),   2000, 104.50],
    ["PAY-001", "Payment", "Euro Supplies GmbH", "EUR", datetime.date(2026,2,10),   8000, 89.50],
    ["PAY-002", "Payment", "Euro Supplies GmbH", "EUR", datetime.date(2026,3,15),  12000, 90.50],
    ["PAY-003", "Payment", "Euro Supplies GmbH", "EUR", datetime.date(2025,12,30),  5000, 88.50],
]
narration = {
    "REC-001": "Part advance received before invoice INV-S-001",
    "REC-002": "2nd receipt against INV-S-001",
    "REC-003": "Final receipt against INV-S-001 (settled over 3 dates)",
    "REC-004": "Single receipt settling two invoices INV-S-002 & INV-S-003",
    "REC-005": "Part receipt against INV-S-004 (balance stays open)",
    "PAY-001": "Part payment against INV-P-001",
    "PAY-002": "Balance payment against INV-P-001 (settled over 2 dates)",
    "PAY-003": "Advance paid before invoice INV-P-002",
}
stl_rows = []
for base in stl_data:
    ref = base[0]
    stl_rows.append(base + [
        "=[@[FC Amount]]*[@[Settlement Rate]]",
        "=SUMIFS(tblAllocations[FC Allocated],tblAllocations[Settlement Ref],[@[Settlement Ref]])",
        "=[@[FC Amount]]-[@[FC Allocated]]",
        '=IF(ROUND([@[FC Unallocated]],2)=0,"Fully Allocated",IF([@[FC Allocated]]=0,"Unallocated","Partly Allocated"))',
        narration[ref],
    ])
last = write_table(ws, 4, stl_columns, stl_rows, "tblSettlements",
                   style="TableStyleMedium6", header_fill=TEAL)
for j in range(8, 12):
    ws.cell(row=4, column=j).fill = fill(SLATE)
r0, r1 = 5, last
fmt_col(ws, "E", FMT_DATE, r0, r1)
fmt_col(ws, "F", FMT_FC, r0, r1)
fmt_col(ws, "G", FMT_RATE, r0, r1)
fmt_col(ws, "H", FMT_AMT, r0, r1)
fmt_col(ws, "I", FMT_FC, r0, r1)
fmt_col(ws, "J", FMT_FC, r0, r1)
stl_widths = [14,15,22,9,15,13,13,18,13,14,15,40]
for j, w in enumerate(stl_widths, start=1):
    ws.column_dimensions[get_column_letter(j)].width = w
ws.freeze_panes = "B5"

# =========================================================================== #
#  5. ALLOCATION & FOREX CALC  (the engine)
# =========================================================================== #
ws = wb.create_sheet("Allocation & Forex Calc")
ws.sheet_view.showGridLines = False
banner(ws, "ALLOCATION & FOREX GAIN / (LOSS) CALCULATION",
        "Link each receipt/payment to the invoice it settles — one line per invoice per settlement date", 20, NAVY)

alloc_columns = [
    "Alloc ID", "Settlement Ref", "Invoice No", "FC Allocated",     # input A-D
    "Transaction Type", "Party", "Currency", "Invoice Date", "Invoice Rate",
    "Settlement Date", "Settlement Rate", "Settlement Type", "Settlement Nature",
    "Base @ Invoice Rate", "Base @ Settlement Rate", "Rate Difference",
    "Forex Gain/(Loss)", "Result", "Days to Settle", "Validation / Remarks",
]
# input allocation lines covering every scenario
alloc_input = [
    ["AL-01", "REC-001", "INV-S-001", 10000],   # advance, invoice over 3 dates
    ["AL-02", "REC-002", "INV-S-001", 12000],
    ["AL-03", "REC-003", "INV-S-001",  8000],
    ["AL-04", "REC-004", "INV-S-002", 12000],   # one receipt -> two invoices
    ["AL-05", "REC-004", "INV-S-003",  8000],
    ["AL-06", "PAY-001", "INV-P-001",  8000],   # purchase paid over 2 dates
    ["AL-07", "PAY-002", "INV-P-001", 12000],
    ["AL-08", "PAY-003", "INV-P-002",  5000],   # advance payment
    ["AL-09", "REC-005", "INV-S-004",  2000],   # partial -> leaves unrealised balance
]
def lk_inv(col):
    return f'=IFERROR(INDEX(tblInvoices[{col}],MATCH([@[Invoice No]],tblInvoices[Invoice No],0)),"")'
def lk_stl(col):
    return f'=IFERROR(INDEX(tblSettlements[{col}],MATCH([@[Settlement Ref]],tblSettlements[Settlement Ref],0)),"")'

alloc_rows = []
for base in alloc_input:
    alloc_rows.append(base + [
        lk_inv("Transaction Type"),
        lk_inv("Party Name"),
        lk_inv("Currency"),
        lk_inv("Invoice Date"),
        lk_inv("Invoice Rate"),
        lk_stl("Settlement Date"),
        lk_stl("Settlement Rate"),
        lk_stl("Settlement Type"),
        '=IF(OR([@[Invoice Date]]="",[@[Settlement Date]]=""),"",IF([@[Settlement Date]]<[@[Invoice Date]],"Advance","On / After Invoice"))',
        '=IF([@[FC Allocated]]="","",[@[FC Allocated]]*[@[Invoice Rate]])',
        '=IF([@[FC Allocated]]="","",[@[FC Allocated]]*[@[Settlement Rate]])',
        '=IF([@[FC Allocated]]="","",[@[Settlement Rate]]-[@[Invoice Rate]])',
        '=IF([@[FC Allocated]]="","",[@[FC Allocated]]*([@[Settlement Rate]]-[@[Invoice Rate]])*IF([@[Transaction Type]]="Sales",1,-1))',
        '=IF([@[Forex Gain/(Loss)]]="","",IF([@[Forex Gain/(Loss)]]>0,"Gain",IF([@[Forex Gain/(Loss)]]<0,"Loss","No Impact")))',
        '=IF(OR([@[Invoice Date]]="",[@[Settlement Date]]=""),"",[@[Settlement Date]]-[@[Invoice Date]])',
        '=IF([@[Invoice No]]="","",IF(IFERROR(INDEX(tblSettlements[Currency],MATCH([@[Settlement Ref]],tblSettlements[Settlement Ref],0)),"")<>[@[Currency]],"Check: currency mismatch",IF([@[FC Allocated]]="","","OK")))',
    ])
last = write_table(ws, 4, alloc_columns, alloc_rows, "tblAllocations",
                   style="TableStyleMedium2", header_fill=BLUE)
# computed headers -> slate; the headline gain/(loss) col -> gold
for j in range(5, 21):
    ws.cell(row=4, column=j).fill = fill(SLATE)
ws.cell(row=4, column=17).fill = fill(GOLD)   # Forex Gain/(Loss)
ws.cell(row=4, column=18).fill = fill(GOLD)   # Result

r0, r1 = 5, last
fmt_col(ws, "D", FMT_FC, r0, r1)
fmt_col(ws, "H", FMT_DATE, r0, r1)
fmt_col(ws, "I", FMT_RATE, r0, r1)
fmt_col(ws, "J", FMT_DATE, r0, r1)
fmt_col(ws, "K", FMT_RATE, r0, r1)
fmt_col(ws, "N", FMT_AMT, r0, r1)
fmt_col(ws, "O", FMT_AMT, r0, r1)
fmt_col(ws, "P", FMT_RATE, r0, r1)
fmt_col(ws, "Q", FMT_AMT, r0, r1)
fmt_col(ws, "S", FMT_INT, r0, r1)
alloc_widths = [9,15,12,13,15,20,9,13,11,15,13,14,16,18,19,13,16,9,12,22]
for j, w in enumerate(alloc_widths, start=1):
    ws.column_dimensions[get_column_letter(j)].width = w
ws.freeze_panes = "E5"

# conditional formatting Forex Gain/(Loss) col Q
rngQ = f"Q{r0}:Q{r1}"
ws.conditional_formatting.add(rngQ, CellIsRule(operator="greaterThan", formula=["0"],
                              font=Font(color=GREEN_F, bold=True), fill=fill(GREEN_BG)))
ws.conditional_formatting.add(rngQ, CellIsRule(operator="lessThan", formula=["0"],
                              font=Font(color=RED_F, bold=True), fill=fill(RED_BG)))
# colour the Settlement Nature 'Advance'
ws.conditional_formatting.add(f"M{r0}:M{r1}",
    CellIsRule(operator="equal", formula=['"Advance"'], font=Font(color=GOLD, bold=True)))

# --------------------------------------------------------------------------- #
#  Data validation (dropdowns) – applied generously so new rows inherit them
# --------------------------------------------------------------------------- #
# Invoices sheet validations
inv_ws = wb["Invoices"]
dv_type = DataValidation(type="list", formula1='"Sales,Purchase"', allow_blank=True)
dv_cur  = DataValidation(type="list", formula1='"USD,EUR,GBP,AED,JPY"', allow_blank=True)
inv_ws.add_data_validation(dv_type); dv_type.add("B5:B1000")
inv_ws.add_data_validation(dv_cur);  dv_cur.add("D5:D1000")

stl_ws = wb["Receipts & Payments"]
dv_stype = DataValidation(type="list", formula1='"Receipt,Payment"', allow_blank=True)
dv_cur2  = DataValidation(type="list", formula1='"USD,EUR,GBP,AED,JPY"', allow_blank=True)
stl_ws.add_data_validation(dv_stype); dv_stype.add("B5:B1000")
stl_ws.add_data_validation(dv_cur2);  dv_cur2.add("D5:D1000")

# Allocation sheet: validate Settlement Ref & Invoice No against the source columns
dv_settref = DataValidation(type="list", formula1="'Receipts & Payments'!$A$5:$A$1000", allow_blank=True)
dv_invno   = DataValidation(type="list", formula1="Invoices!$A$5:$A$1000", allow_blank=True)
ws.add_data_validation(dv_settref); dv_settref.add("B5:B1000")
ws.add_data_validation(dv_invno);   dv_invno.add("C5:C1000")

# =========================================================================== #
#  6. DASHBOARD
# =========================================================================== #
ws = wb.create_sheet("Dashboard")
ws.sheet_view.showGridLines = False
banner(ws, "FOREX GAIN / (LOSS) DASHBOARD",
        "Auto-summary of realised, unrealised and net foreign-exchange impact", 6, NAVY)

def kpi_block(ws, top, title, color, rows):
    ws.merge_cells(start_row=top, start_column=2, end_row=top, end_column=4)
    c = ws.cell(row=top, column=2, value=title)
    c.font = Font(bold=True, color=WHITE, size=11)
    c.fill = fill(color)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[top].height = 20
    rr = top + 1
    for label, formula, fmt, is_gl in rows:
        l = ws.cell(row=rr, column=2, value=label)
        l.font = cell_font()
        l.fill = fill(BAND)
        l.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        l.border = BORDER
        ws.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=3)
        v = ws.cell(row=rr, column=4, value=formula)
        v.number_format = fmt
        v.font = cell_font(bold=True)
        v.alignment = RIGHT
        v.border = BORDER
        if is_gl:
            v.font = cell_font(bold=True)
        rr += 1
    return rr

# Realised
r = kpi_block(ws, 4, "REALISED FOREX (money actually moved)", TEAL, [
    ("Realised Gains",  '=SUMIF(tblAllocations[Forex Gain/(Loss)],">0")', FMT_AMT, True),
    ("Realised Losses", '=SUMIF(tblAllocations[Forex Gain/(Loss)],"<0")', FMT_AMT, True),
    ("Net Realised Gain/(Loss)", "=SUM(tblAllocations[Forex Gain/(Loss)])", FMT_AMT, True),
    ("  • on Sales (receipts)",   '=SUMIFS(tblAllocations[Forex Gain/(Loss)],tblAllocations[Transaction Type],"Sales")', FMT_AMT, True),
    ("  • on Purchases (payments)", '=SUMIFS(tblAllocations[Forex Gain/(Loss)],tblAllocations[Transaction Type],"Purchase")', FMT_AMT, True),
])
r += 1
# Unrealised
r = kpi_block(ws, r, "UNREALISED FOREX (open balance revalued at closing rate)", GOLD, [
    ("Net Unrealised Gain/(Loss)", "=SUM(tblInvoices[Unrealised Gain/(Loss)])", FMT_AMT, True),
    ("  • on Sales (debtors open)",   '=SUMIFS(tblInvoices[Unrealised Gain/(Loss)],tblInvoices[Transaction Type],"Sales")', FMT_AMT, True),
    ("  • on Purchases (creditors open)", '=SUMIFS(tblInvoices[Unrealised Gain/(Loss)],tblInvoices[Transaction Type],"Purchase")', FMT_AMT, True),
])
r += 1
# Net total
r = kpi_block(ws, r, "TOTAL FOREX IMPACT TO P&L", NAVY, [
    ("Net Forex Gain/(Loss) (Realised + Unrealised)",
     "=SUM(tblAllocations[Forex Gain/(Loss)])+SUM(tblInvoices[Unrealised Gain/(Loss)])", FMT_AMT, True),
])
r += 1
# Counters
r = kpi_block(ws, r, "DATA OVERVIEW", SLATE, [
    ("Invoices captured",          "=COUNTA(tblInvoices[Invoice No])", FMT_INT, False),
    ("Receipts & Payments captured", "=COUNTA(tblSettlements[Settlement Ref])", FMT_INT, False),
    ("Allocation lines",           "=COUNTA(tblAllocations[Alloc ID])", FMT_INT, False),
    ("Advance settlements",        '=COUNTIF(tblAllocations[Settlement Nature],"Advance")', FMT_INT, False),
    ("Invoices fully settled",     '=COUNTIF(tblInvoices[Settlement Status],"Fully Settled")', FMT_INT, False),
    ("Invoices part/open",         '=COUNTIFS(tblInvoices[Settlement Status],"<>Fully Settled")', FMT_INT, False),
])

# conditional format the gain/loss value cells (column D) where they are amounts
for rr in range(5, r):
    cell = ws.cell(row=rr, column=4)
    if cell.number_format == FMT_AMT:
        addr = f"D{rr}"
        ws.conditional_formatting.add(addr, CellIsRule(operator="greaterThan", formula=["0"], font=Font(color=GREEN_F, bold=True)))
        ws.conditional_formatting.add(addr, CellIsRule(operator="lessThan", formula=["0"], font=Font(color=RED_F, bold=True)))

ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 30
ws.column_dimensions["C"].width = 14
ws.column_dimensions["D"].width = 20
ws.column_dimensions["E"].width = 3

# Negative numbers shown in brackets red for all amount formats -> apply a P&L style number format
PNL_FMT = '#,##0.00;[Red](#,##0.00)'
for sheet, cols, r0, r1 in [
    (wb["Invoices"], ["M","N","O"], 5, 5+len(inv_rows)-1),
    (wb["Allocation & Forex Calc"], ["Q","P"], 5, 5+len(alloc_rows)-1),
]:
    for col in cols:
        for rr in range(r0, r1+1):
            sheet[f"{col}{rr}"].number_format = PNL_FMT
for rr in range(5, r):
    if ws.cell(row=rr, column=4).number_format == FMT_AMT:
        ws.cell(row=rr, column=4).number_format = PNL_FMT

# tab colours
wb["Read Me"].sheet_properties.tabColor = NAVY
wb["Settings"].sheet_properties.tabColor = SLATE
wb["Invoices"].sheet_properties.tabColor = BLUE
wb["Receipts & Payments"].sheet_properties.tabColor = TEAL
wb["Allocation & Forex Calc"].sheet_properties.tabColor = GOLD
wb["Dashboard"].sheet_properties.tabColor = NAVY

wb.calculation.fullCalcOnLoad = True

out = "Foreign_Gain_Loss_Workbook.xlsx"
wb.save(out)
print("Saved", out)
