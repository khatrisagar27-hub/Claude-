"""Build the birthday-party receipt register workbook from the handwritten pages.

Source: two photographed pages of the handwritten register for the
"Ria & Sejal" birthday party dated 23-08-2026 (Gujarati script, red ink).
Entries 1-28 are on page 1, entries 29-59 on page 2.

Every figure below is transcribed from those photos - nothing is inferred.
Rows whose handwriting could not be read with confidence carry "Yes" in the
Verify column and must be checked against the book before the sheet is used.
"""

import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = "Receipt_Details_Sejal_Ria.xlsx"

S, R, C = "sejal", "ria", "common"

# sr, name (transliterated), name (Gujarati as written), place / note, page,
# {column: amount}, verify?, remark
ROWS = [
    (1,  "Kiran Somabhai Vyas", "કિરણ સોમાભાઈ વ્યાસ", "—", 1, {}, "No", "Casserole set given to Sejal — gift in kind, no cash amount"),
    (2,  "Kishorbhai Hiralal Patel", "કિશોરભાઈ હીરાલાલ પટેલ", "Andhavas / sankada", 1, {R: 5100}, "No", "Marked 'રાયુ' (Ria)"),
    (3,  "Jinal Prajapati", "જીનલ પ્રજાપતિ", "—", 1, {S: 500}, "No", "Marked 'સેજલને'"),
    (4,  "Aman Shah / Manasi Soni", "અમન શાહ, માનસી સોની", "—", 1, {C: 1000}, "Yes", "First name unclear"),
    (5,  "Dhaval Mukeshbhai Jadhav", "ધવલ મુકેશભાઈ જાધવ", "40 Sugandha row, Maravad", 1, {C: 501}, "Yes", ""),
    (6,  "Pravin Rajput (Kiran)", "પ્રવીણ રાજપૂત (કીરણ)", "Kirtan Kusela", 1, {C: 501}, "Yes", ""),
    (7,  "Lakshmanbhai Agrawal", "લક્ષ્મણભાઈ અગ્રવાલ", "Sejal kaka", 1, {R: 1100}, "No", "Marked 'રાયુ'"),
    (8,  "Lakshmanbhai Agrawal", "લક્ષ્મણભાઈ અગ્રવાલ", "Sejal kaka (ditto)", 1, {S: 1100}, "No", ""),
    (9,  "Flower bouquet (Kirti)", "ફૂલોના બુકે (કીર્તિ)", "—", 1, {R: 500}, "No", ""),
    (10, "Flower bouquet (Kirti)", "ફૂલોના બુકે (કીર્તિ)", "ditto", 1, {S: 500}, "No", ""),
    (11, "Ketan Dhaneshbhai Agrawal", "કેતન ધનેશભાઈ અગ્રવાલ", "Sangar kaka", 1, {R: 1100}, "No", ""),
    (12, "Ketan Dhaneshbhai Agrawal", "કેતન ધનેશભાઈ અગ્રવાલ", "ditto", 1, {S: 1100}, "No", ""),
    (13, "Salamsinh Agrawal", "સલમસિંહ અગ્રવાલ", "Sejal kaka", 1, {R: 1100}, "Yes", "Name reading uncertain"),
    (14, "Salamsinh Agrawal", "સલમસિંહ અગ્રવાલ", "ditto", 1, {S: 1100}, "Yes", ""),
    (15, "Nilam Dhanendra Gupta", "નીલમ ધનેન્દ્ર ગુપ્તા", "—", 1, {S: 500, R: 500}, "No", "Braced in the book as 500 Sejal + 500 Ria"),
    (16, "Kadri Gaurav Agrawal", "કાદ્રી ગૌરવ અગ્રવાલ", "—", 1, {C: 500}, "Yes", "No child named against this entry"),
    (17, "Ravi Shubhamkumar Agrawal", "રવિ શુભમકુમાર અગ્રવાલ", "—", 1, {S: 500}, "No", ""),
    (18, "Neha Shubhamkumar Agrawal", "નેહા શુભમકુમાર અગ્રવાલ", "ditto", 1, {R: 500}, "No", ""),
    (19, "Niya Pritesh Agrawal", "નીયા પ્રીતેશ અગ્રવાલ", "—", 1, {S: 500}, "No", ""),
    (20, "Niya Pritesh Agrawal", "નીયા પ્રીતેશ અગ્રવાલ", "ditto", 1, {R: 500}, "No", ""),
    (21, "Keyur Agrawal", "કેયૂર અગ્રવાલ", "—", 1, {S: 1100}, "No", ""),
    (22, "Keyur Agrawal", "કેયૂર અગ્રવાલ", "ditto", 1, {R: 1100}, "No", ""),
    (23, "Shakeshbhai Agrawal", "શાકેશભાઈ અગ્રવાલ", "Anjali nayapa", 1, {C: 701}, "Yes", ""),
    (24, "Rahul Kanjibhai Patel", "રાહુલ કાનજીભાઈ પટેલ", "Gam Gorasa", 1, {C: 700}, "Yes", ""),
    (25, "Somaben Vyas", "સોમાબેન વ્યાસ", "Ambalkaram", 1, {C: 500}, "No", ""),
    (26, "Nilaben Lad", "નીલાબેન લાડ", "Ambalkaram", 1, {C: 500}, "No", ""),
    (27, "Pramilaben Patel", "પ્રમીલાબેન પટેલ", "Ambalkaram", 1, {C: 500}, "No", ""),
    (28, "Unnatiben Narendrakumar Patel", "ઉન્નતીબેન નરેન્દ્રકુમાર પટેલ", "Ahmedabad", 1, {C: 500}, "Yes", ""),

    (29, "Ochhiben Paragkumar Patel", "ઓછીબેન પરાગકુમાર પટેલ", "Ahmedabad", 2, {C: 500}, "Yes", ""),
    (30, "Kaushalyaben Hudiya", "કૌશલ્યાબેન હુડીયા", "Ambalkaram", 2, {C: 500}, "Yes", ""),
    (31, "Shariben Pandya", "શારીબેન પંડ્યા", "Ambalkaram", 2, {C: 500}, "Yes", ""),
    (32, "Ushaben Joshi", "ઉષાબેન જોષી", "Ambalkaram", 2, {C: 501}, "No", ""),
    (33, "Kanjibhai Darji", "કાનજીભાઈ દરજી", "Ambalkaram", 2, {C: 501}, "No", ""),
    (34, "Ashokbhai Soni", "અશોકભાઈ સોની", "Ambalkaram", 2, {C: 500}, "Yes", ""),
    (35, "Vadekabhai Panchal", "વાડેકાભાઈ પંચાલ", "Ambalkaram", 2, {C: 500}, "Yes", ""),
    (36, "Shaileshbhai Chavda", "શીલેશભાઈ ચાવડા", "Ambalkaram", 2, {C: 500}, "No", ""),
    (37, "Nilavbhai Darji", "નીળવભાઈ દરજી", "Ambalkaram", 2, {C: 1001}, "Yes", ""),
    (38, "Dhanvanarayan Mistri", "ધનવનારાયણ મીસ્ત્રી", "Kirtan Kumar A/1", 2, {R: 501}, "No", "Marked 'રાયા'"),
    (39, "Ajaybhai Pancholi", "અજયભાઈ પંચોલી", "Kirtan Kumesh A/2", 2, {R: 501}, "No", "Marked 'રાયા'"),
    (40, "Gamma Thakor Dhira", "ગામમા ઠાકોર ધીરા", "Kom Dhulabh", 2, {C: 501}, "Yes", "A struck-out name follows in the book"),
    (41, "Raman Koli", "રમણ કોળી", "A/3 Kirtan ku. Mahesh", 2, {C: 500}, "Yes", ""),
    (42, "Sanjilaben Prajapati", "સંજીલાબેન પ્રજાપતિ", "Rayanagar", 2, {C: 501}, "Yes", ""),
    (43, "Anubhai Bhaichandbhai", "અનુભાઈ ભાઈચંદભાઈ", "Davaliphali", 2, {S: 1251}, "Yes", "Marked 'સેજલ'"),
    (44, "Anandkumar", "અનંદકુમાર", "—", 2, {C: 751}, "Yes", "Name and marking both unclear"),
    (45, "Dileshbhai", "દિલેશભાઈ", "Ahmedabad", 2, {C: 1001}, "Yes", ""),
    (46, "Divyesh Hiralal Patel", "દિવ્યેશ હીરાલાલ પટેલ", "Vivekanand Society, Ahmedabad", 2, {C: 1000}, "Yes", ""),
    (47, "Divyesh Hiralal Patel", "દિવ્યેશ હીરાલાલ પટેલ", "ditto", 2, {S: 500}, "Yes", "Marked 'સેજલ'"),
    (48, "Vave Ramgambhai", "વાવે રામગમભાઈ", "—", 2, {C: 501}, "Yes", ""),
    (49, "Rameshbhai Agrawal", "રમેશભાઈ અગ્રવાલ", "Agropalu", 2, {S: 2500}, "Yes", "Marked 'સેજલ'"),
    (51, "Kalpana Rakeshkumar Patel", "કલ્પના રાકેશકુમાર પટેલ", "Ahmedabad", 2, {R: 2500}, "Yes", "Marked 'રાયા'"),
    (52, "Kalpana Rakeshkumar Patel", "કલ્પના રાકેશકુમાર પટેલ", "Ahmedabad (ditto)", 2, {S: 1101}, "Yes", "Marked 'સેજલને'"),
    (53, "Neha Vrushang Patel", "નેહા વૃષંગ પટેલ", "—", 2, {R: 1101}, "Yes", "Marked 'રાયાને'"),
    (54, "Neha Vrushang Patel", "નેહા વૃષંગ પટેલ", "Ahmedabad", 2, {S: 701}, "Yes", "Marked 'સેજલને'"),
    (55, "Neha Vrushang Patel", "નેહા વૃષંગ પટેલ", "Ahmedabad (ditto)", 2, {R: 701}, "Yes", "Marked 'રાયાને'"),
    (56, "Jagadbhai Rajpam", "જગદભાઈ રાજપમ", "Kirtan ku. m.", 2, {S: 200}, "Yes", "Marked 'સેજલને'"),
    (57, "Lagnanagar Kapar", "લગ્નનગર કપર", "—", 2, {S: 300}, "Yes", "Marked 'સેજલને'"),
    (58, "Lagnanagar Kapar", "લગ્નનગર કપર", "ditto", 2, {C: 500}, "Yes", ""),
    (59, "Lagnanagar Kapar", "લગ્નનગર કપર", "ditto", 2, {C: 500}, "Yes", ""),
    ("", "Vishal Patel", "વિશાલ પટેલ", "Kendre — written below the last numbered entry", 2, {C: 15000}, "Yes", "Unnumbered line at the foot of page 2"),
]

# Totals written by hand in the register itself, for reconciliation.
HAND_PAGE1 = 24303
HAND_GRAND = 61985   # bottom-right of page 2; partly cut off in the photo

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
BODY = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
MONEY = '#,##0;(#,##0);-'
SEJAL_FILL = PatternFill("solid", fgColor="FCE4EC")
RIA_FILL = PatternFill("solid", fgColor="E3F2FD")
TOTAL_FILL = PatternFill("solid", fgColor="FFF2CC")
FLAG_FILL = PatternFill("solid", fgColor="FFF9C4")

wb = Workbook()

# ---------------------------------------------------------------- register --
ws = wb.active
ws.title = "Receipt Details"

ws["A1"] = "Ria & Sejal — Birthday Party Receipt Register"
ws["A1"].font = Font(name="Arial", size=14, bold=True, color="1F3864")
ws["A2"] = "Date of party: 23-08-2026 · transcribed from the handwritten register (2 pages) · all amounts in ₹"
ws["A2"].font = Font(name="Arial", size=9, italic=True, color="595959")

HEADERS = ["Sr. No.", "Name of giver", "Name (as written)", "Place / reference",
           "Page", "Sejal (₹)", "Ria (₹)", "Common / not specified (₹)",
           "Total (₹)", "Verify?", "Remarks"]
for col, head in enumerate(HEADERS, start=1):
    cell = ws.cell(row=4, column=col, value=head)
    cell.font = HDR_FONT
    cell.fill = HDR_FILL
    cell.border = BORDER
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

first = 5
for i, (sr, en, gu, place, page, amts, verify, remark) in enumerate(ROWS):
    r = first + i
    ws.cell(row=r, column=1, value=sr)
    ws.cell(row=r, column=2, value=en)
    ws.cell(row=r, column=3, value=gu)
    ws.cell(row=r, column=4, value=place)
    ws.cell(row=r, column=5, value=page)
    ws.cell(row=r, column=6, value=amts.get(S))
    ws.cell(row=r, column=7, value=amts.get(R))
    ws.cell(row=r, column=8, value=amts.get(C))
    ws.cell(row=r, column=9, value=f"=SUM(F{r}:H{r})")
    ws.cell(row=r, column=10, value=verify)
    ws.cell(row=r, column=11, value=remark)
    for col in range(1, 12):
        c = ws.cell(row=r, column=col)
        c.font = BODY
        c.border = BORDER
        if col in (6, 7, 8, 9):
            c.number_format = MONEY
            c.alignment = Alignment(horizontal="right")
        elif col in (1, 5, 10):
            c.alignment = Alignment(horizontal="center")
        else:
            c.alignment = Alignment(horizontal="left", vertical="center")
    ws.cell(row=r, column=6).fill = SEJAL_FILL
    ws.cell(row=r, column=7).fill = RIA_FILL
    ws.cell(row=r, column=9).fill = TOTAL_FILL
    if verify == "Yes":
        ws.cell(row=r, column=10).fill = FLAG_FILL

last = first + len(ROWS) - 1
tot = last + 1
ws.cell(row=tot, column=2, value="TOTAL")
for col in (6, 7, 8, 9):
    L = get_column_letter(col)
    ws.cell(row=tot, column=col, value=f"=SUM({L}{first}:{L}{last})")
for col in range(1, 12):
    c = ws.cell(row=tot, column=col)
    c.font = BOLD
    c.border = Border(left=THIN, right=THIN, top=Side(style="double"), bottom=Side(style="double"))
    c.fill = TOTAL_FILL
    if col in (6, 7, 8, 9):
        c.number_format = MONEY
        c.alignment = Alignment(horizontal="right")

for col, width in zip("ABCDEFGHIJK", (8, 30, 30, 30, 6, 13, 13, 22, 13, 9, 46)):
    ws.column_dimensions[col].width = width
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A4:K{last}"
ws.sheet_view.showGridLines = False

note = ws.cell(row=tot + 2, column=2,
               value="Sejal / Ria columns follow the child named against that line in the register "
                     "('સેજલ' / 'રાયા'). Lines with no child named are put under Common. "
                     "Rows flagged Verify = Yes were hard to read in the photograph — check them against the book.")
note.font = Font(name="Arial", size=9, italic=True, color="595959")
note.alignment = Alignment(vertical="top", wrap_text=True)
ws.merge_cells(start_row=tot + 2, start_column=2, end_row=tot + 4, end_column=11)

# ----------------------------------------------------------------- summary --
sm = wb.create_sheet("Summary")
sm["A1"] = "Bifurcation Summary"
sm["A1"].font = Font(name="Arial", size=14, bold=True, color="1F3864")
sm["A2"] = "Every figure below is a formula over the 'Receipt Details' sheet."
sm["A2"].font = Font(name="Arial", size=9, italic=True, color="595959")

DET = "'Receipt Details'"
blocks = [
    ("A4", "Head", "Amount (₹)", "No. of entries"),
    ("A5", "Sejal", f"=SUM({DET}!F{first}:F{last})", f"=COUNT({DET}!F{first}:F{last})"),
    ("A6", "Ria", f"=SUM({DET}!G{first}:G{last})", f"=COUNT({DET}!G{first}:G{last})"),
    ("A7", "Common / not specified", f"=SUM({DET}!H{first}:H{last})", f"=COUNT({DET}!H{first}:H{last})"),
    ("A8", "Grand total", f"=SUM(B5:B7)", f"=SUM(C5:C7)"),
]
for addr, label, amount, count in blocks:
    r = int(addr[1:])
    sm.cell(row=r, column=1, value=label)
    sm.cell(row=r, column=2, value=amount)
    sm.cell(row=r, column=3, value=count)

for r in range(4, 9):
    for col in range(1, 4):
        c = sm.cell(row=r, column=col)
        c.border = BORDER
        c.font = BOLD if r in (4, 8) else BODY
        if r == 4:
            c.fill = HDR_FILL
            c.font = HDR_FONT
            c.alignment = Alignment(horizontal="center")
        elif col >= 2:
            c.number_format = MONEY if col == 2 else "0"
            c.alignment = Alignment(horizontal="right")
sm["B5"].fill = SEJAL_FILL
sm["B6"].fill = RIA_FILL
sm["B8"].fill = TOTAL_FILL
sm["C8"].fill = TOTAL_FILL

sm["A10"] = "If the common / unmarked amount is shared equally"
sm["A10"].font = BOLD
sm["A11"] = "Sejal (own + half of common)"
sm["B11"] = "=B5+B7/2"
sm["A12"] = "Ria (own + half of common)"
sm["B12"] = "=B6+B7/2"
for r in (11, 12):
    sm.cell(row=r, column=1).font = BODY
    sm.cell(row=r, column=2).font = BODY
    sm.cell(row=r, column=2).number_format = MONEY
    for col in (1, 2):
        sm.cell(row=r, column=col).border = BORDER
sm["B11"].fill = SEJAL_FILL
sm["B12"].fill = RIA_FILL

sm["A14"] = "Reconciliation with the totals written in the register"
sm["A14"].font = BOLD
recon = [
    ("Page 1 total as per this sheet", f"=SUMIF({DET}!E{first}:E{last},1,{DET}!I{first}:I{last})"),
    ("Page 1 total written in the register", HAND_PAGE1),
    ("Difference (page 1)", "=B15-B16"),
    ("Grand total as per this sheet", "=B8"),
    ("Grand total written in the register", HAND_GRAND),
    ("Difference (grand total)", "=B18-B19"),
]
for i, (label, val) in enumerate(recon):
    r = 15 + i
    sm.cell(row=r, column=1, value=label).font = BODY
    c = sm.cell(row=r, column=2, value=val)
    c.font = BODY
    c.number_format = MONEY
    for col in (1, 2):
        sm.cell(row=r, column=col).border = BORDER
sm["B17"].font = BOLD
sm["B20"].font = BOLD

sm["A22"] = ("Page 1 reconciles exactly. The grand total at the foot of page 2 is partly cut off in the "
             "photograph and was read as 61,985 — the difference above is the amount still to be traced, "
             "and the rows flagged Verify = Yes on the detail sheet are where it will be.")
sm["A22"].font = Font(name="Arial", size=9, italic=True, color="C00000")
sm["A22"].alignment = Alignment(vertical="top", wrap_text=True)
sm.merge_cells("A22:D25")

for col, width in zip("ABCD", (42, 18, 16, 16)):
    sm.column_dimensions[col].width = width
sm.sheet_view.showGridLines = False

# -------------------------------------------------------------------- notes --
nt = wb.create_sheet("Notes")
nt["A1"] = "How this sheet was prepared"
nt["A1"].font = Font(name="Arial", size=14, bold=True, color="1F3864")
lines = [
    "Source: two photographs of the handwritten register for the Ria & Sejal birthday party, dated 23-08-2026.",
    "Page 1 carries entries 1 to 28 and a page total of 24,303. Page 2 carries that figure forward and runs to entry 59.",
    "",
    "Column rules",
    "  Sejal — the line is marked 'સેજલ' / 'સેજલને' in the register.",
    "  Ria — the line is marked 'રાયા' / 'રાયુ' / 'રાયાને' (written as Raya; recorded here as Ria).",
    "  Common / not specified — no child is named on that line in the register.",
    "  Total — a formula, Sejal + Ria + Common for that row.",
    "",
    "Entry 1 (Kiran Somabhai Vyas) is a casserole set given to Sejal — a gift in kind, so it carries no amount.",
    "Entry 15 (Nilam Dhanendra Gupta) is written with a brace in the book: 500 to Sejal and 500 to Ria.",
    "The register's own numbering skips 50; the numbers here follow the book exactly.",
    "The last line of page 2 (Vishal Patel, 15,000) is unnumbered in the register.",
    "",
    "Names are transliterated from Gujarati handwriting and the original reading is kept alongside in column C.",
    "Rows flagged Verify = Yes could not be read with confidence — correct the name or the amount from the book.",
    "Amounts are whole rupees; the register writes them as '500=00' style, i.e. rupees and paise.",
]
for i, line in enumerate(lines):
    c = nt.cell(row=3 + i, column=1, value=line)
    c.font = BOLD if line in ("Column rules",) else Font(name="Arial", size=10)
nt.column_dimensions["A"].width = 120
nt.sheet_view.showGridLines = False

wb.save(OUT)
print("wrote", OUT)


# ---------------------------------------------------------------------------
# LibreOffice is unavailable here, and openpyxl writes formulas with no cached
# result - which reads back as blank in pandas and in most previewers. So the
# same arithmetic is done in Python and injected as the cached <v> for each
# formula cell. The formulas themselves are untouched and Excel recalculates
# them on open.
# ---------------------------------------------------------------------------
import re
import shutil
import zipfile

sejal_col = [r[5].get(S) for r in ROWS]
ria_col = [r[5].get(R) for r in ROWS]
common_col = [r[5].get(C) for r in ROWS]
pages = [r[4] for r in ROWS]
row_totals = [sum(v for v in (a, b, c) if v) for a, b, c in zip(sejal_col, ria_col, common_col)]

sum_s = sum(v for v in sejal_col if v)
sum_r = sum(v for v in ria_col if v)
sum_c = sum(v for v in common_col if v)
cnt_s = sum(1 for v in sejal_col if v is not None)
cnt_r = sum(1 for v in ria_col if v is not None)
cnt_c = sum(1 for v in common_col if v is not None)
page1 = sum(t for t, p in zip(row_totals, pages) if p == 1)
grand = sum_s + sum_r + sum_c

detail_values = {f"I{first + i}": row_totals[i] for i in range(len(ROWS))}
detail_values.update({f"F{tot}": sum_s, f"G{tot}": sum_r, f"H{tot}": sum_c, f"I{tot}": grand})

summary_values = {
    "B5": sum_s, "B6": sum_r, "B7": sum_c, "B8": grand,
    "C5": cnt_s, "C6": cnt_r, "C7": cnt_c, "C8": cnt_s + cnt_r + cnt_c,
    "B11": sum_s + sum_c / 2, "B12": sum_r + sum_c / 2,
    "B15": page1, "B16": HAND_PAGE1, "B17": page1 - HAND_PAGE1,
    "B18": grand, "B19": HAND_GRAND, "B20": grand - HAND_GRAND,
}


# openpyxl already emits an empty <v/> placeholder after each <f>. A cell may
# carry only one <v>, so that placeholder is replaced - appending a second one
# produces a file Excel refuses to open.
_FORMULA_CELL = re.compile(
    r'(<c r="([A-Z]+\d+)"[^>]*>\s*<f>[^<]*</f>)\s*(?:<v\s*/>|<v>[^<]*</v>)?'
)


def _inject(xml: str, values: dict) -> str:
    def repl(match):
        head, ref = match.group(1), match.group(2)
        if ref not in values:
            return match.group(0)
        val = values[ref]
        val = int(val) if float(val).is_integer() else val
        return f"{head}<v>{val}</v>"

    return _FORMULA_CELL.sub(repl, xml)


tmp = OUT + ".tmp"
shutil.copy(OUT, tmp)
with zipfile.ZipFile(tmp) as src, zipfile.ZipFile(OUT, "w", zipfile.ZIP_DEFLATED) as dst:
    for item in src.infolist():
        data = src.read(item.filename)
        if item.filename == "xl/worksheets/sheet1.xml":
            data = _inject(data.decode("utf-8"), detail_values).encode("utf-8")
        elif item.filename == "xl/worksheets/sheet2.xml":
            data = _inject(data.decode("utf-8"), summary_values).encode("utf-8")
        dst.writestr(item, data)
os.remove(tmp)

print(f"Sejal {sum_s:,} | Ria {sum_r:,} | Common {sum_c:,} | Grand {grand:,} | page 1 {page1:,}")
